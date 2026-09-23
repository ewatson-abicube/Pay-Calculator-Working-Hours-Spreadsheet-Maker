import csv
import json
import os
import re
import sys
import tkinter as tk
from datetime import datetime
from tkinter import ttk

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None

def get_base_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

BASE_DIR = get_base_dir()
DATA_FILE = os.path.join(BASE_DIR, "weekly_payroll_data.json")
CSV_EXPORT_PATH = os.path.join(BASE_DIR, "weekly_payroll.csv")
XLSX_EXPORT_PATH = os.path.join(BASE_DIR, "weekly_payroll.xlsx")
EXPORT_DATA_FILE = os.path.join(BASE_DIR, "export_rows_data.json")


def validate_non_negative_input(value):
    if value == "":
        return True

    try:
        numeric_value = float(value)
    except ValueError:
        return False

    return numeric_value >= 0


def attach_numeric_validation(entry):
    entry.configure(validate="key", validatecommand=(entry.register(validate_non_negative_input), "%P"))


def validate_non_negative_hours(value):
    if value == "":
        return True

    try:
        numeric_value = float(value)
    except ValueError:
        return False

    return numeric_value >= 0


def validate_text_only(value):
    if value == "":
        return True

    return bool(re.fullmatch(r"[A-Za-z\s'.-]*", value))


def sanitize_numeric_value(value):
    if value == "":
        return ""

    try:
        numeric_value = float(value)
    except ValueError:
        return ""

    return "0" if numeric_value < 0 else str(numeric_value)


def save_inputs():
    data = {
        "rate": sanitize_numeric_value(rate_entry.get()),
        "days": {}
    }

    for day, entry in day_entries.items():
        data["days"][day] = sanitize_numeric_value(entry.get())

    with open(DATA_FILE, "w") as file:
        json.dump(data, file)


def load_inputs():
    if not os.path.exists(DATA_FILE):
        return

    try:
        with open(DATA_FILE, "r") as file:
            data = json.load(file)
    except (json.JSONDecodeError, OSError):
        return

    if "rate" in data:
        rate_value = sanitize_numeric_value(data["rate"])
        rate_entry.delete(0, tk.END)
        rate_entry.insert(0, rate_value if rate_value != "" else "0")

    if "days" in data:
        for day, value in data["days"].items():
            if day in day_entries:
                day_entry_value = sanitize_numeric_value(value)
                day_entries[day].delete(0, tk.END)
                day_entries[day].insert(0, day_entry_value if day_entry_value != "" else "0")


def get_payroll_row():
    try:
        rate = float(rate_entry.get() or 0)
    except ValueError:
        return None

    total_hours = 0.0

    for entry in day_entries.values():
        total_hours += float(entry.get() or 0)

    return {
        "Date": datetime.now().strftime("%Y-%m-%d"),
        "Rate": rate,
        "Mon": float(day_entries["Mon"].get() or 0),
        "Tue": float(day_entries["Tue"].get() or 0),
        "Wed": float(day_entries["Wed"].get() or 0),
        "Thu": float(day_entries["Thu"].get() or 0),
        "Fri": float(day_entries["Fri"].get() or 0),
        "Sat": float(day_entries["Sat"].get() or 0),
        "Sun": float(day_entries["Sun"].get() or 0),
        "Total Hours": total_hours,
        "Total Pay": total_hours * rate
    }


def calculate_total():
    try:
        rate = float(rate_entry.get())
        total_hours = 0.0

        for day_name, entry in day_entries.items():
            hours = float(entry.get() or 0)
            total_hours += hours

        total = total_hours * rate
        result_var.set(f"Total Pay: ${total:,.2f}")
        hours_total_var.set(f"Total Hours: {total_hours:,.2f} hrs")
        save_inputs()
    except ValueError:
        result_var.set("Please enter valid numbers")
        hours_total_var.set("")
        save_inputs()


def reset_calculator():
    rate_entry.delete(0, tk.END)
    rate_entry.insert(0, "15")

    for entry in day_entries.values():
        entry.delete(0, tk.END)

    result_var.set("Total Pay: $0.00")
    hours_total_var.set("Total Hours: 0.00 hrs")
    save_inputs()


def export_row_key(date_value, hours_value, staff_value, task_value):
    return (
        str(date_value).strip(),
        str(hours_value).strip(),
        str(staff_value).strip(),
        str(task_value).strip(),
    )


def row_exists_in_export_tree(date_value, hours_value, staff_value, task_value):
    target_key = export_row_key(date_value, hours_value, staff_value, task_value)

    for item in export_tree.get_children():
        values = export_tree.item(item, "values")
        if len(values) != 4:
            continue

        if export_row_key(*values[:4]) == target_key:
            return True

    return False


def save_export_rows():
    rows = []
    for item in export_tree.get_children():
        values = export_tree.item(item, "values")
        if len(values) != 4:
            continue
        rows.append({
            "Date": values[0],
            "Hours": values[1],
            "Staff": values[2],
            "Task": values[3]
        })

    with open(EXPORT_DATA_FILE, "w") as file:
        json.dump(rows, file)


def load_export_rows():
    if not os.path.exists(EXPORT_DATA_FILE):
        return

    try:
        with open(EXPORT_DATA_FILE, "r") as file:
            rows = json.load(file)
    except (json.JSONDecodeError, OSError):
        return

    for row in rows:
        export_tree.insert("", "end", values=(row.get("Date", ""), row.get("Hours", ""), row.get("Staff", ""), row.get("Task", "")))


def add_export_row():
    date_value = date_entry.get().strip()
    hours_value = hours_entry.get().strip()
    staff_value = staff_entry.get().strip()
    task_value = task_entry.get().strip()

    if not date_value or not hours_value or not staff_value or not task_value:
        return

    try:
        hours_numeric = float(hours_value)
    except ValueError:
        return

    if hours_numeric < 0:
        return

    if row_exists_in_export_tree(date_value, hours_value, staff_value, task_value):
        date_entry.delete(0, tk.END)
        date_entry.insert(0, datetime.now().strftime("%m/%d/%Y"))
        hours_entry.delete(0, tk.END)
        staff_entry.delete(0, tk.END)
        task_entry.delete(0, tk.END)
        return

    export_tree.insert("", "end", values=(date_value, hours_value, staff_value, task_value))
    save_export_rows()

    date_entry.delete(0, tk.END)
    date_entry.insert(0, datetime.now().strftime("%m/%d/%Y"))
    hours_entry.delete(0, tk.END)
    staff_entry.delete(0, tk.END)
    task_entry.delete(0, tk.END)


def delete_selected_export_row(event=None):
    selected_items = export_tree.selection()
    if not selected_items:
        return

    for item in selected_items:
        export_tree.delete(item)

    save_export_rows()


def get_existing_export_row_keys():
    existing_keys = set()

    if os.path.exists(CSV_EXPORT_PATH):
        with open(CSV_EXPORT_PATH, newline="") as csv_file:
            reader = csv.DictReader(csv_file)
            for row in reader:
                if not row:
                    continue
                date_value = row.get("Date", "")
                hours_value = row.get("Hours", "")
                staff_value = row.get("Staff", "")
                task_value = row.get("Task", "")
                existing_keys.add(export_row_key(date_value, hours_value, staff_value, task_value))

    if Workbook is not None and os.path.exists(XLSX_EXPORT_PATH):
        wb = load_workbook(XLSX_EXPORT_PATH)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 4:
                continue
            date_value, hours_value, staff_value, task_value = row[:4]
            existing_keys.add(export_row_key(date_value, hours_value, staff_value, task_value))

    return existing_keys


def export_export_rows():
    rows = []
    seen_keys = set()
    existing_keys = get_existing_export_row_keys()

    for item in export_tree.get_children():
        values = export_tree.item(item, "values")
        if len(values) != 4:
            continue

        date_value, hours_value, staff_value, task_value = values
        row_key = export_row_key(date_value, hours_value, staff_value, task_value)
        if row_key in seen_keys or row_key in existing_keys:
            continue

        seen_keys.add(row_key)
        rows.append({
            "Date": date_value,
            "Hours": hours_value,
            "Staff": staff_value,
            "Task": task_value
        })

    if not rows:
        return

    csv_exists = os.path.exists(CSV_EXPORT_PATH)
    with open(CSV_EXPORT_PATH, "a", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=["Date", "Hours", "Staff", "Task"])
        if not csv_exists:
            writer.writeheader()
        for row in rows:
            writer.writerow(row)

    if Workbook is not None:
        if os.path.exists(XLSX_EXPORT_PATH):
            wb = load_workbook(XLSX_EXPORT_PATH)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Export"
            ws.append(["Date", "Hours", "Staff", "Task"])

        for row in rows:
            ws.append([row["Date"], row["Hours"], row["Staff"], row["Task"]])

        wb.save(XLSX_EXPORT_PATH)

    save_export_rows()

    try:
        if os.path.exists(XLSX_EXPORT_PATH):
            os.startfile(XLSX_EXPORT_PATH)
        else:
            os.startfile(CSV_EXPORT_PATH)
    except Exception:
        pass


def clear_export_table():
    for row in export_tree.get_children():
        export_tree.delete(row)

    if os.path.exists(CSV_EXPORT_PATH):
        os.remove(CSV_EXPORT_PATH)

    if Workbook is not None:
        if os.path.exists(XLSX_EXPORT_PATH):
            os.remove(XLSX_EXPORT_PATH)

    if os.path.exists(EXPORT_DATA_FILE):
        os.remove(EXPORT_DATA_FILE)


def on_close():
    save_inputs()
    save_export_rows()
    root.destroy()


root = tk.Tk()
root.title("Weekly Payroll Calculator")
root.geometry("1100x720")
root.minsize(1000, 650)
root.configure(bg="#eef4ff")

style = ttk.Style(root)
style.theme_use("clam")
style.configure("TFrame", background="#FFFFFF")
style.configure("TLabel", background="#eef4ff", foreground="#000000", font=("Segoe UI", 10))
style.configure("Header.TLabel", font=("Segoe UI", 24, "bold"), foreground="#111827")
style.configure("Section.TLabel", font=("Segoe UI", 11, "bold"), foreground="#1d4ed8")
style.configure("TEntry", padding=8, fieldbackground="#ffffff")
style.configure("Accent.TButton", font=("Segoe UI", 11, "bold"), foreground="#ffffff", background="#2563eb")
style.map("Accent.TButton", background=[("active", "#1d4ed8")])
style.configure("Secondary.TButton", font=("Segoe UI", 11, "bold"), foreground="#0f172a", background="#dbeafe")
style.map("Secondary.TButton", background=[("active", "#bfdbfe")])
style.configure("SectionDivider.TSeparator", background="#1d4ed8")
style.configure("SectionDivider.TSeparator", lightcolor="#1d4ed8", darkcolor="#1d4ed8", troughcolor="#e2e8f0", thickness=8)
style.configure("Black.TSeparator", background="#000000")

main_frame = ttk.Frame(root, padding=25)
main_frame.pack(fill="both", expand=True)

header = ttk.Label(main_frame, text="Weekly Work Pay Calculator", style="Header.TLabel")
header.pack(pady=(0, 12))

calculator_frame = ttk.Frame(main_frame)
calculator_frame.pack(fill="x", pady=(0, 18))

calculator_label = ttk.Label(calculator_frame, text="Calculator", style="Section.TLabel")
calculator_label.pack(anchor="w", pady=(0, 8))

inputs_frame = ttk.Frame(calculator_frame)
inputs_frame.pack(fill="x", pady=(0, 12))

rate_label = ttk.Label(inputs_frame, text="Pay per Hour:")
rate_label.grid(row=0, column=0, sticky="w", padx=(0, 10), pady=8)

rate_entry = ttk.Entry(inputs_frame, width=18)
rate_entry.grid(row=0, column=1, sticky="w")
rate_entry.insert(0, "15")
attach_numeric_validation(rate_entry)

inputs_frame.columnconfigure(1, weight=1)

weekdays = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

daily_frame = ttk.Frame(calculator_frame)
daily_frame.pack(fill="x", pady=8)

day_entries = {}
for index, day in enumerate(weekdays):
    card = tk.Frame(daily_frame, padx=10, pady=8, bg="#ffffff")
    card.grid(row=0, column=index, padx=8, pady=4, sticky="nsew")

    day_label = ttk.Label(card, text=day, font=("Segoe UI", 11, "bold"), foreground="#0f172a")
    day_label.pack(anchor="center", pady=(0, 6))

    entry = ttk.Entry(card, width=8, justify="center")
    entry.pack(anchor="center")
    attach_numeric_validation(entry)
    day_entries[day] = entry

    card.grid_columnconfigure(0, weight=1)

daily_frame.columnconfigure(tuple(range(len(weekdays))), weight=1)

button_row = ttk.Frame(calculator_frame)
button_row.pack(fill="x", pady=(18, 12))

calculate_button = ttk.Button(button_row, text="Calculate Total", style="Accent.TButton", command=calculate_total)
calculate_button.pack(side="left", padx=(0, 10))

reset_button = ttk.Button(button_row, text="Reset", style="Secondary.TButton", command=reset_calculator)
reset_button.pack(side="left")

result_var = tk.StringVar(value="Total Pay: $0.00")
hours_total_var = tk.StringVar(value="Total Hours: 0.00 hrs")

result_label = ttk.Label(calculator_frame, textvariable=result_var, font=("Segoe UI", 18, "bold"), foreground="#0f766e")
result_label.pack(pady=(0, 6))

hours_label = ttk.Label(calculator_frame, textvariable=hours_total_var, font=("Segoe UI", 10, "bold"), foreground="#334155")
hours_label.pack(pady=(0, 8))

separator = ttk.Separator(main_frame, orient="horizontal", style="SectionDivider.TSeparator")
separator.pack(fill="x", pady=(12, 14))

export_section = ttk.Frame(main_frame)
export_section.pack(fill="both", expand=True, pady=(0, 0))

export_title = ttk.Label(export_section, text="Spreadsheet Export", style="Section.TLabel")
export_title.pack(anchor="w", pady=(0, 10))

export_form = ttk.Frame(export_section)
export_form.pack(fill="x", pady=(0, 10))

ttk.Label(export_form, text="Date:").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=5)
date_entry = ttk.Entry(export_form, width=18)
date_entry.grid(row=0, column=1, sticky="w", padx=(0, 12), pady=5)
date_entry.insert(0, datetime.now().strftime("%m/%d/%Y"))

ttk.Label(export_form, text="Hours:").grid(row=0, column=2, sticky="w", padx=(0, 8), pady=5)
hours_entry = ttk.Entry(export_form, width=12)
hours_entry.grid(row=0, column=3, sticky="w", padx=(0, 12), pady=5)
hours_entry.configure(validate="key", validatecommand=(hours_entry.register(validate_non_negative_hours), "%P"))

ttk.Label(export_form, text="Staff:").grid(row=0, column=4, sticky="w", padx=(0, 8), pady=5)
staff_entry = ttk.Entry(export_form, width=18)
staff_entry.grid(row=0, column=5, sticky="w", padx=(0, 12), pady=5)
staff_entry.configure(validate="key", validatecommand=(staff_entry.register(validate_text_only), "%P"))

ttk.Label(export_form, text="Task:").grid(row=0, column=6, sticky="w", padx=(0, 8), pady=5)
task_entry = ttk.Entry(export_form, width=22)
task_entry.grid(row=0, column=7, sticky="w", pady=5)
task_entry.configure(validate="key", validatecommand=(task_entry.register(validate_text_only), "%P"))

add_row_button = ttk.Button(export_form, text="Add Row", style="Secondary.TButton", command=add_export_row)
add_row_button.grid(row=1, column=0, sticky="w", pady=(10, 0), columnspan=2)

delete_row_button = ttk.Button(export_form, text="Delete Selected", style="Secondary.TButton", command=delete_selected_export_row)
delete_row_button.grid(row=1, column=2, sticky="w", pady=(10, 0), padx=(0, 10), columnspan=2)

export_button = ttk.Button(export_form, text="Export Spreadsheet", style="Accent.TButton", command=export_export_rows)
export_button.grid(row=1, column=4, sticky="w", pady=(10, 0), padx=(0, 10), columnspan=2)

clear_log_button = ttk.Button(export_form, text="Clear Spreadsheet Log", style="Secondary.TButton", command=clear_export_table)
clear_log_button.grid(row=1, column=6, sticky="w", pady=(10, 0), columnspan=2)

export_tree = ttk.Treeview(export_section, columns=("Date", "Hours", "Staff", "Task"), show="headings", height=12)
export_tree.heading("Date", text="Date")
export_tree.heading("Hours", text="Hours")
export_tree.heading("Staff", text="Staff")
export_tree.heading("Task", text="Task")
export_tree.column("Date", width=120, anchor="center")
export_tree.column("Hours", width=100, anchor="center")
export_tree.column("Staff", width=180, anchor="center")
export_tree.column("Task", width=240, anchor="center")
export_tree.pack(fill="both", expand=True)

export_tree.bind("<Double-Button-1>", lambda event: edit_tree_cell(event))
export_tree.bind("<Delete>", delete_selected_export_row)


def edit_tree_cell(event):
    item = export_tree.identify_row(event.y)
    if not item:
        return

    column_id = export_tree.identify_column(event.x)
    column_index = int(column_id.replace("#", "")) - 1
    if column_index < 0 or column_index > 3:
        return

    column_name = export_tree["columns"][column_index]
    current_value = export_tree.item(item, "values")[column_index]

    x, y, width, height = export_tree.bbox(item, column_id)
    entry = tk.Entry(export_tree, show="", bd=0, relief="flat", highlightthickness=1, highlightbackground="#bdbdbd", highlightcolor="#1d4ed8", background="#ffffff", foreground="#000000", insertbackground="#000000", font=("Segoe UI", 10))
    entry.place(x=x, y=y, width=width, height=height)
    entry.insert(0, current_value)
    entry.selection_range(0, tk.END)
    entry.focus_set()

    def save_edit(event=None):
        new_value = entry.get().strip()
        export_tree.set(item, column_name, new_value)
        entry.destroy()
        save_export_rows()

    entry.bind("<Return>", save_edit)
    entry.bind("<FocusOut>", save_edit)


if not os.path.exists(EXPORT_DATA_FILE):
    sample_rows = [
        ("9/21/2026", "3", "Eric Watson", "Task 1"),
        ("9/22/2026", "4", "Eric Watson", "Task 2"),
        ("9/23/2026", "2", "Eric Watson", "Task 3"),
    ]
    for row in sample_rows:
        export_tree.insert("", "end", values=row)
    save_export_rows()
else:
    load_export_rows()

root.protocol("WM_DELETE_WINDOW", on_close)
load_inputs()
rate_entry.focus()
root.mainloop()